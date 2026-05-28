#!/usr/bin/env python3
"""Smart Interview workbook writer + validator.

The ONLY entry point that touches .xlsx files. Three sub-commands:
  seed   — create workbook with canonical schema, populate Guide + Assumption Matrix
  write  — append scored data into the remaining sheets
  qa     — validate schema, coverage, MVP purity, no-leakage; writes hooks/.qa_state

Schema is frozen here. Any drift aborts with non-zero exit.
"""
from __future__ import annotations
import argparse, json, sys, datetime
from pathlib import Path

try:
    from openpyxl import Workbook, load_workbook
except ImportError:
    sys.exit("ERROR: openpyxl not installed. Run: python3 -m pip install --user openpyxl")

SCHEMA = {
    "Guide":              ["question_id", "section", "text", "assumption_ids"],
    "Assumption Matrix":  ["id", "text", "source", "score", "class", "evidence_refs"],
    "Scoping Matrix":     ["use_case_id", "description", "priority", "linked_assumptions"],
    "MVP Specifications": ["feature_id", "name", "description", "linked_assumptions"],
    "Evidence Log":       ["timestamp", "assumption_id", "quote", "line_ref", "polarity", "score"],
    "Open Questions":     ["id", "question", "reason", "linked_assumptions"],
}
VALID_CLASSES = {"STRONG", "CONTESTED", "WEAK", "INVALIDATED"}
QA_STATE = Path(__file__).resolve().parent.parent / "hooks" / ".qa_state"


def _check_schema(wb) -> list[str]:
    errs = []
    if set(wb.sheetnames) != set(SCHEMA):
        errs.append(f"sheet set mismatch: have {wb.sheetnames}, want {list(SCHEMA)}")
        return errs
    for sheet, cols in SCHEMA.items():
        row1 = [c.value for c in wb[sheet][1]] if wb[sheet].max_row else []
        if row1[: len(cols)] != cols:
            errs.append(f"{sheet}: header drift — have {row1}, want {cols}")
    return errs


def _safe_cell(v):
    if v is None:
        return ""
    s = str(v)
    if "\n" in s:
        s = s.replace("\n", " ")
    return s


def cmd_seed(args):
    out = Path(args.path)
    out.parent.mkdir(parents=True, exist_ok=True)
    wb = Workbook()
    wb.remove(wb.active)
    for sheet, cols in SCHEMA.items():
        ws = wb.create_sheet(sheet)
        ws.append(cols)

    guide = json.loads(Path(args.guide).read_text()) if args.guide else []
    for q in guide:
        wb["Guide"].append([
            _safe_cell(q.get("question_id")),
            _safe_cell(q.get("section")),
            _safe_cell(q.get("text")),
            _safe_cell(",".join(q.get("assumption_ids", []))),
        ])

    assumptions = json.loads(Path(args.assumptions).read_text()) if args.assumptions else []
    for a in assumptions:
        wb["Assumption Matrix"].append([
            _safe_cell(a.get("id")),
            _safe_cell(a.get("text")),
            _safe_cell(a.get("source")),
            "", "", "",
        ])
    wb.save(out)
    print(json.dumps({"seeded": str(out), "guide_rows": len(guide), "assumption_rows": len(assumptions)}))


def cmd_write(args):
    wb = load_workbook(args.path)
    errs = _check_schema(wb)
    if errs:
        sys.exit("SCHEMA ABORT: " + "; ".join(errs))

    payload = json.loads(Path(args.payload).read_text())
    written = {}

    # Assumption Matrix: rewrite class/score/evidence per assumption id (in place).
    if "assumption_matrix" in payload:
        idx = {row[0].value: row for row in wb["Assumption Matrix"].iter_rows(min_row=2) if row[0].value}
        for entry in payload["assumption_matrix"]:
            r = idx.get(entry["id"])
            if not r:
                sys.exit(f"unknown assumption id {entry['id']}")
            if entry.get("class") not in VALID_CLASSES:
                sys.exit(f"invalid class {entry.get('class')} for {entry['id']}")
            r[3].value = entry.get("score", "")
            r[4].value = entry["class"]
            r[5].value = _safe_cell(entry.get("evidence_refs", ""))
        written["Assumption Matrix"] = len(payload["assumption_matrix"])

    for sheet_key, sheet in [
        ("mvp_specifications", "MVP Specifications"),
        ("scoping_matrix",     "Scoping Matrix"),
        ("open_questions",     "Open Questions"),
    ]:
        rows = payload.get(sheet_key, [])
        for row in rows:
            wb[sheet].append([_safe_cell(row.get(c)) for c in SCHEMA[sheet]])
        if rows:
            written[sheet] = len(rows)

    # Evidence Log: append-only.
    ev = payload.get("evidence_log", [])
    ts = datetime.datetime.utcnow().isoformat(timespec="seconds") + "Z"
    for s in ev:
        wb["Evidence Log"].append([
            ts,
            _safe_cell(s.get("assumption_id")),
            _safe_cell(s.get("quote")),
            _safe_cell(s.get("line_ref")),
            s.get("polarity", 0),
            s.get("score", ""),
        ])
    if ev:
        written["Evidence Log"] = len(ev)

    wb.save(args.path)
    print(json.dumps({"sheets_written": list(written), "rows_per_sheet": written}))


def cmd_qa(args):
    wb = load_workbook(args.path)
    failures = _check_schema(wb)

    am_rows = list(wb["Assumption Matrix"].iter_rows(min_row=2, values_only=True))
    am_rows = [r for r in am_rows if r and r[0]]
    am = {r[0]: {"class": r[4], "evidence": r[5]} for r in am_rows}

    ev_ids = {r[1] for r in wb["Evidence Log"].iter_rows(min_row=2, values_only=True) if r and r[1]}
    for aid in am:
        if aid not in ev_ids:
            failures.append(f"coverage: {aid} missing from Evidence Log")
        if not am[aid]["evidence"]:
            failures.append(f"evidence: {aid} has empty evidence_refs")

    for r in wb["MVP Specifications"].iter_rows(min_row=2, values_only=True):
        if not r or not r[0]:
            continue
        linked = (r[3] or "").split(",")
        for lid in [x.strip() for x in linked if x.strip()]:
            cls = am.get(lid, {}).get("class")
            if cls != "STRONG":
                failures.append(f"MVP purity: row {r[0]} links {lid} class={cls}")

    for r in wb["Scoping Matrix"].iter_rows(min_row=2, values_only=True):
        if not r or not r[0]:
            continue
        linked = (r[3] or "").split(",")
        for lid in [x.strip() for x in linked if x.strip()]:
            cls = am.get(lid, {}).get("class")
            if cls in {"WEAK", "INVALIDATED"}:
                failures.append(f"no-leakage: Scoping {r[0]} links {lid} class={cls}")

    contested = {aid for aid, v in am.items() if v["class"] == "CONTESTED"}
    oq_links = set()
    for r in wb["Open Questions"].iter_rows(min_row=2, values_only=True):
        if r and r[3]:
            for x in str(r[3]).split(","):
                oq_links.add(x.strip())
    for aid in contested - oq_links:
        failures.append(f"open-questions: CONTESTED {aid} not present in Open Questions")

    QA_STATE.parent.mkdir(parents=True, exist_ok=True)
    if failures:
        QA_STATE.write_text("fail " + " | ".join(failures))
        print(json.dumps({"pass": False, "failures": failures}, indent=2))
        sys.exit(1)
    QA_STATE.write_text("pass " + datetime.datetime.utcnow().isoformat())
    print(json.dumps({"pass": True, "failures": []}))


def main():
    p = argparse.ArgumentParser()
    sub = p.add_subparsers(dest="cmd", required=True)

    s = sub.add_parser("seed");  s.add_argument("--path", required=True); s.add_argument("--guide"); s.add_argument("--assumptions"); s.set_defaults(fn=cmd_seed)
    s = sub.add_parser("write"); s.add_argument("--path", required=True); s.add_argument("--payload", required=True); s.set_defaults(fn=cmd_write)
    s = sub.add_parser("qa");    s.add_argument("--path", required=True); s.set_defaults(fn=cmd_qa)

    args = p.parse_args()
    args.fn(args)


if __name__ == "__main__":
    main()
