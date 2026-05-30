#!/usr/bin/env python3
"""PowerPoint readout deck from the study ledger (v0.2).

A deck is a synthesis VIEW, like the Word PRD — built from the SQLite store, never a
new source of data. ~8 slides:
  1. Title            5. Open questions (CONTESTED)
  2. Method + saturation  6. Prioritisation (table)
  3. Verdict summary  7. Where we're still thin (next steps)
  4. Validated scope (STRONG, with a key quote each)
"""
from __future__ import annotations
from pathlib import Path


def _bullets(slide, frame_text_pairs):
    """frame_text_pairs: list of (text, level)."""
    tf = slide.placeholders[1].text_frame
    tf.clear()
    first = True
    for text, level in frame_text_pairs:
        p = tf.paragraphs[0] if first else tf.add_paragraph()
        p.text = text
        p.level = level
        first = False


def build_pptx(path, study, scores, assumptions, interviews, signals, ranked, sat):
    from pptx import Presentation
    from pptx.util import Inches, Pt

    prs = Presentation()
    TITLE, TITLE_CONTENT, TITLE_ONLY = prs.slide_layouts[0], prs.slide_layouts[1], prs.slide_layouts[5]

    def _pos_quote(aid):
        for s in signals:
            if s["assumption_id"] == aid and s["score"] > 0 and s.get("quote"):
                return f"“{s['quote']}” — {s['interview_id']}"
        return ""

    # 1. Title
    s = prs.slides.add_slide(TITLE)
    s.shapes.title.text = f"{study['name']}"
    s.placeholders[1].text = f"User research readout · {len(interviews)} interview(s)"

    # 2. Method + saturation
    s = prs.slides.add_slide(TITLE_CONTENT)
    s.shapes.title.text = "Method"
    consented = sum(1 for iv in interviews if iv.get("consent"))
    segs = sorted({iv.get("segment") for iv in interviews if iv.get("segment")})
    _bullets(s, [
        (f"{len(interviews)} interviews; consent on record: {consented}/{len(interviews)}", 0),
        (f"Segments: {', '.join(segs) or 'n/a'}", 0),
        (f"Saturation: {sat['verdict']}", 0),
        ("Verdict computed in code (strength / consensus / confidence); every claim traces to a quote.", 0),
    ])

    # 3. Verdict summary
    s = prs.slides.add_slide(TITLE_CONTENT)
    s.shapes.title.text = "Assumption verdicts"
    counts = {}
    for v in scores.values():
        counts[v["class"]] = counts.get(v["class"], 0) + 1
    _bullets(s, [(f"{k}: {counts.get(k, 0)}", 0)
                 for k in ("STRONG", "CONTESTED", "WEAK", "INVALIDATED")])

    # 4. Validated scope (STRONG)
    s = prs.slides.add_slide(TITLE_CONTENT)
    s.shapes.title.text = "Validated scope (STRONG)"
    strong = [a for a in assumptions if scores[a["id"]]["class"] == "STRONG"]
    pairs = []
    if not strong:
        pairs.append(("No assumption is STRONG with adequate confidence yet.", 0))
    for a in strong:
        pairs.append((f"{a['id']}  {a['text']}", 0))
        q = _pos_quote(a["id"])
        if q:
            pairs.append((q, 1))
    _bullets(s, pairs)

    # 5. Open questions (CONTESTED)
    s = prs.slides.add_slide(TITLE_CONTENT)
    s.shapes.title.text = "Open questions (CONTESTED)"
    contested = [a for a in assumptions if scores[a["id"]]["class"] == "CONTESTED"]
    _bullets(s, [(f"{a['id']}  {a['text']}", 0) for a in contested]
                 or [("None.", 0)])

    # 6. Prioritisation (table)
    s = prs.slides.add_slide(TITLE_ONLY)
    s.shapes.title.text = "Prioritisation (weighted)"
    rows = (ranked or [])[:8]
    tbl = s.shapes.add_table(len(rows) + 1, 4, Inches(0.5), Inches(1.6),
                             Inches(9), Inches(0.4 * (len(rows) + 1))).table
    for j, h in enumerate(("Use Case", "Score", "Priority", "Linked")):
        tbl.cell(0, j).text = h
    for i, u in enumerate(rows, 1):
        tbl.cell(i, 0).text = u["name"]
        tbl.cell(i, 1).text = str(u["weighted_score"])
        tbl.cell(i, 2).text = u["priority"]
        tbl.cell(i, 3).text = u.get("linked_assumptions") or ""
    if not rows:
        tbl.cell(1, 0).text = "No use cases scored."

    # 7. Where we're still thin
    s = prs.slides.add_slide(TITLE_CONTENT)
    s.shapes.title.text = "Next steps"
    nxt = []
    if sat["thin"]:
        nxt.append((f"Still thin (n<=2), need more interviews: {', '.join(sat['thin'])}", 0))
    if sat["needs_more"]:
        nxt.append((f"Positive but unconfirmed: {', '.join(sat['needs_more'])}", 0))
    if contested:
        nxt.append((f"Resolve contested: {', '.join(a['id'] for a in contested)}", 0))
    if not nxt:
        nxt.append(("Coverage looks saturated — proceed to PRD.", 0))
    _bullets(s, nxt)

    Path(path).parent.mkdir(parents=True, exist_ok=True)
    prs.save(path)
    return path
